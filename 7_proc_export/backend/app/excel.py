# -*- coding: utf-8 -*-
"""把查詢結果寫成單一 Excel，清單／內容／明細／簽核名單各一個 tab。

用 write_only 模式：逐欄展開後一列可能上百欄，兩萬列在一般模式下會把
記憶體吃光。write_only 一次只留一列在記憶體裡。

欄位標題一律用「欄位名稱」；名稱查不到（表單定義沒填標籤）時退回欄位 id，
不拿 id 混充中文名稱。同一個 tab 內名稱重複時附上 id 區別，避免兩欄同名。
表格控件的明細不塞進單頭的一格 —— 每個表格自成一個 tab，一列一筆。
"""

import datetime
import io
import re

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Excel 不接受的控制字元；不濾掉會在寫入時整份炸掉
ILLEGAL_RE = re.compile(r'[\000-\010\013\014\016-\037]')
CELL_LIMIT = 32767
SHEET_NAME_LIMIT = 31

DATE_FORMAT = 'yyyy/mm/dd hh:mm:ss'

HEADER_FILL = PatternFill('solid', fgColor='DDEBF7')
HEADER_FONT = Font(bold=True)
HEADER_ALIGN = Alignment(vertical='center', wrap_text=True)

# 清單 tab 的欄位（基本單頭 + 狀態時程 + 作廢原因）
LIST_COLUMNS = [
    ('processSerialNumber', '流程單號', 26),
    ('formSerialNumber', '表單單號', 22),
    ('processId', '流程 ID', 26),
    ('processName', '流程名稱', 22),
    ('subject', '主旨', 34),
    ('requesterId', '申請人工號', 12),
    ('requesterName', '申請人姓名', 12),
    ('orgUnitId', '申請部門代號', 12),
    ('orgUnitName', '申請部門名稱', 20),
    ('createdTime', '申請日期', 19),
    ('stateName', '目前狀態', 10),
    ('closedTime', '結案時間', 19),
    ('elapsedDays', '流程歷時（天）', 13),
    ('currentStep', '目前關卡', 18),
    ('currentPerformer', '目前待辦人', 18),
    ('completedSteps', '已簽關卡數', 11),
    ('abortComment', '作廢／終止原因', 28),
    ('abortedBy', '作廢／終止人', 14),
    ('parseError', '表單解析狀況', 22),
]

SIGN_COLUMNS = [
    ('processSerialNumber', '流程單號', 26),
    ('formSerialNumber', '表單單號', 22),
    ('subject', '主旨', 30),
    ('stepIndex', '關卡序', 8),
    ('workItemName', '關卡名稱', 20),
    ('performerId', '簽核人工號', 12),
    ('performerName', '簽核人姓名', 12),
    ('createdTime', '收件時間', 19),
    ('completedTime', '完成時間', 19),
    ('elapsedDays', '停留天數', 10),
    ('stateName', '關卡狀態', 10),
    ('comment', '簽核意見', 40),
]

CONTENT_FIXED = [
    ('processSerialNumber', '流程單號', 26),
    ('formSerialNumber', '表單單號', 22),
    ('subject', '主旨', 30),
    ('requesterId', '申請人工號', 12),
    ('requesterName', '申請人姓名', 12),
    ('createdTime', '申請日期', 19),
    ('stateName', '目前狀態', 10),
]

DETAIL_FIXED = [
    ('processSerialNumber', '流程單號', 26),
    ('formSerialNumber', '表單單號', 22),
    ('rowIndex', '明細列序', 10),
]


def _clean_text(value):
    """字串化並濾掉 Excel 拒收的控制字元，過長就截斷並註明。"""
    if value is None:
        return ''
    if isinstance(value, (int, float, datetime.datetime, datetime.date)):
        return value
    text = ILLEGAL_RE.sub('', str(value))
    if len(text) > CELL_LIMIT:
        return text[:CELL_LIMIT - 20] + '…（內容過長截斷）'
    return text


def _elapsed_days(start, end):
    if not start or not end:
        return ''
    return round((end - start).total_seconds() / 86400.0, 2)


def _unique_headers(pairs):
    """(id, 名稱) → 標題清單。名稱重複時附上 id，避免兩欄同名分不出來。"""
    counts = {}
    for _, name in pairs:
        counts[name] = counts.get(name, 0) + 1
    headers = []
    for field_id, name in pairs:
        headers.append('%s (%s)' % (name, field_id) if counts[name] > 1 else name)
    return headers


def _safe_sheet_name(name, used):
    """Excel 工作表名稱不得超過 31 字且不可重複，也不可含 []:*?/\\"""
    cleaned = re.sub(r'[\[\]:*?/\\]', '_', name).strip() or '工作表'
    cleaned = cleaned[:SHEET_NAME_LIMIT]
    candidate = cleaned
    suffix = 2
    while candidate in used:
        tail = '_%d' % suffix
        candidate = cleaned[:SHEET_NAME_LIMIT - len(tail)] + tail
        suffix += 1
    used.add(candidate)
    return candidate


def _write_header(sheet, headers, widths):
    cells = []
    for text in headers:
        cell = WriteOnlyCell(sheet, value=text)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cells.append(cell)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'A2'
    sheet.append(cells)


def _row_cells(sheet, values):
    cells = []
    for value in values:
        value = _clean_text(value)
        cell = WriteOnlyCell(sheet, value=value)
        if isinstance(value, (datetime.datetime, datetime.date)):
            cell.number_format = DATE_FORMAT
        cells.append(cell)
    return cells


# ---------------------------------------------------------------- 各個 tab

def _sheet_list(workbook, used_names, records):
    sheet = workbook.create_sheet(_safe_sheet_name('清單', used_names))
    _write_header(sheet, [c[1] for c in LIST_COLUMNS], [c[2] for c in LIST_COLUMNS])
    for record in records:
        row = dict(record['summary'])
        row.update(record['progress'])
        row['elapsedDays'] = _elapsed_days(row.get('createdTime'), row.get('closedTime'))
        sheet.append(_row_cells(sheet, [row.get(key, '') for key, _, _ in LIST_COLUMNS]))
    return sheet


def _sheet_content(workbook, used_names, records, selected_fields):
    """selected_fields: [(欄位id, 欄位名稱)]，順序即欄序。"""
    sheet = workbook.create_sheet(_safe_sheet_name('內容', used_names))
    headers = [c[1] for c in CONTENT_FIXED] + _unique_headers(selected_fields)
    widths = [c[2] for c in CONTENT_FIXED] + [18] * len(selected_fields)
    _write_header(sheet, headers, widths)

    for record in records:
        summary = record['summary']
        parsed = record['parsed']
        values = [summary.get(key, '') for key, _, _ in CONTENT_FIXED]
        for field_id, _ in selected_fields:
            values.append(parsed['fields'].get(field_id, ''))
        sheet.append(_row_cells(sheet, values))
    return sheet


def _sheet_details(workbook, used_names, records, grids):
    """每個表格控件一個 tab。grids: [{'id','name','columns':[{'id','name'}]}]。"""
    sheets = []
    for grid in grids:
        columns = grid.get('columns') or []
        if not columns:
            continue
        title = '明細-%s' % (grid.get('name') or grid['id'])
        sheet = workbook.create_sheet(_safe_sheet_name(title, used_names))
        headers = ([c[1] for c in DETAIL_FIXED] +
                   _unique_headers([(c['id'], c['name']) for c in columns]))
        widths = [c[2] for c in DETAIL_FIXED] + [18] * len(columns)
        _write_header(sheet, headers, widths)

        for record in records:
            summary = record['summary']
            rows = record['parsed']['grids'].get(grid['id']) or []
            for index, grid_row in enumerate(rows, 1):
                values = [summary.get('processSerialNumber', ''),
                          summary.get('formSerialNumber', ''), index]
                for column in columns:
                    values.append(grid_row.get(column['id'], ''))
                sheet.append(_row_cells(sheet, values))
        sheets.append(sheet)
    return sheets


def _sheet_signatures(workbook, used_names, records):
    sheet = workbook.create_sheet(_safe_sheet_name('簽核名單', used_names))
    _write_header(sheet, [c[1] for c in SIGN_COLUMNS], [c[2] for c in SIGN_COLUMNS])
    for record in records:
        summary = record['summary']
        for index, item in enumerate(record['items'], 1):
            row = dict(item)
            row['processSerialNumber'] = summary.get('processSerialNumber', '')
            row['formSerialNumber'] = summary.get('formSerialNumber', '')
            row['subject'] = summary.get('subject', '')
            row['stepIndex'] = index
            row['elapsedDays'] = _elapsed_days(item.get('createdTime'),
                                               item.get('completedTime'))
            sheet.append(_row_cells(sheet, [row.get(key, '') for key, _, _ in SIGN_COLUMNS]))
    return sheet


def _sheet_field_map(workbook, used_names, selected_fields, grids, catalog_fields):
    """欄位對照：稽核要能回推「這一欄的原始欄位 id 是什麼、名稱哪裡來的」。"""
    sheet = workbook.create_sheet(_safe_sheet_name('欄位對照', used_names))
    _write_header(sheet, ['所在工作表', '欄位名稱', '欄位 ID', '型別', '名稱來源'],
                  [14, 26, 30, 14, 26])
    source_of = {f['id']: f for f in catalog_fields}
    for field_id, name in selected_fields:
        meta = source_of.get(field_id) or {}
        origin = '表單定義的標籤' if meta.get('named') else '表單定義未填標籤，顯示欄位 ID'
        sheet.append(_row_cells(sheet, ['內容', name, field_id,
                                        meta.get('type', ''), origin]))
    for grid in grids:
        for column in grid.get('columns') or []:
            origin = ('表格欄位標題（caption）' if column.get('named')
                      else '表格定義未填標題，顯示欄位 ID')
            sheet.append(_row_cells(sheet, ['明細-%s' % (grid.get('name') or grid['id']),
                                            column['name'], column['id'], '', origin]))
    return sheet


def _sheet_summary(workbook, used_names, meta):
    """匯出條件與範圍：稽核文件必須能說清楚「這份是怎麼撈出來的」。"""
    sheet = workbook.create_sheet(_safe_sheet_name('匯出條件', used_names))
    _write_header(sheet, ['項目', '內容'], [22, 90])
    for label, value in meta:
        sheet.append(_row_cells(sheet, [label, value]))
    return sheet


# ---------------------------------------------------------------- 對外入口

def build(records, options):
    """組出 Excel，回傳 bytes。

    records: [{'summary':…, 'parsed':…, 'items':[…], 'progress':…}]
    options: {'list','content','signatures'} 布林、selected_fields、grids、
             catalog_fields、meta（匯出條件說明）
    """
    workbook = Workbook(write_only=True)
    used_names = set()

    _sheet_summary(workbook, used_names, options.get('meta') or [])
    if options.get('list'):
        _sheet_list(workbook, used_names, records)
    if options.get('content'):
        selected = options.get('selected_fields') or []
        _sheet_content(workbook, used_names, records, selected)
        _sheet_details(workbook, used_names, records, options.get('grids') or [])
        _sheet_field_map(workbook, used_names, selected, options.get('grids') or [],
                         options.get('catalog_fields') or [])
    if options.get('signatures'):
        _sheet_signatures(workbook, used_names, records)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
