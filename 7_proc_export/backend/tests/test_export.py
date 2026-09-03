# -*- coding: utf-8 -*-
"""不需要資料庫的單元測試：XML 解析、欄位標題、Excel 產出、查詢條件。

會連資料庫的部分（catalog / query）不放進這裡 —— 那要有 .env 與內網才跑得動，
測試不該因為換一台機器就紅掉。資料庫那條路徑靠 README 的實測步驟驗證。
"""

import datetime
import io
import os
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import catalog, excel, form_xml, query  # noqa: E402


# ---------------------------------------------------------------- 表單 XML

SAMPLE = (
    '<ITAccount_Application>'
    '<s_Req id="s_Req" dataType="java.lang.String">S014060</s_Req>'
    '<Dlg id="Dlg" label="蔣佳緯" dataType="java.lang.String">S112009</Dlg>'
    '<Empty id="Empty"></Empty>'
    '<Grid11 id="Grid11"><records>'
    '<record id="Grid11_0"><item id="G_a">EasyFlow</item><item id="G_b">申請</item></record>'
    '<record id="Grid11_1"><item id="G_a">Internet</item></record>'
    '</records></Grid11>'
    '</ITAccount_Application>'
)


def test_parse_splits_header_and_grid():
    parsed = form_xml.parse(SAMPLE)
    assert parsed['formId'] == 'ITAccount_Application'
    assert parsed['fields']['s_Req'] == 'S014060'
    assert parsed['order'] == ['s_Req', 'Dlg', 'Empty']
    # 表格不能混進單頭欄位，否則會被當成一般欄位塞進一個儲存格
    assert 'Grid11' not in parsed['fields']
    assert len(parsed['grids']['Grid11']) == 2
    assert parsed['grids']['Grid11'][0]['G_b'] == '申請'


def test_display_value_keeps_both_label_and_code():
    """DIALOGINPUT 類欄位的內文是工號、label 是姓名，兩者都要留住。"""
    parsed = form_xml.parse(SAMPLE)
    assert parsed['fields']['Dlg'] == '蔣佳緯 (S112009)'


def test_unwrap_removes_form_collection_wrapper():
    wrapped = ('<com.dsc.nana.services.webservice.FormCollection>'
               '<fieldValues>%s</fieldValues>'
               '</com.dsc.nana.services.webservice.FormCollection>'
               % SAMPLE.replace('<', '&lt;').replace('>', '&gt;'))
    parsed = form_xml.parse(wrapped)
    assert parsed['wrapped'] == 1
    assert parsed['formId'] == 'ITAccount_Application'


def test_safe_parse_reports_error_instead_of_raising():
    result = form_xml.safe_parse('<not closed')
    assert result['error']
    assert result['fields'] == {}


# ---------------------------------------------------------------- 表格欄位標題

def test_grid_columns_reads_caption():
    form_xml_text = (
        '<com.dsc.nana.domain.form.ListElementDefinition>'
        '<listItems class="list">'
        '<com.dsc.nana.domain.form.ListItem>'
        '<itemOrder>0</itemOrder><id>G_s_Sys_Name</id><binding>s_Sys_Name</binding>'
        '<elementDefinition class="x"><id>ignored</id></elementDefinition>'
        '<caption>系統名稱</caption><hidden>false</hidden>'
        '</com.dsc.nana.domain.form.ListItem>'
        '</listItems>'
        '</com.dsc.nana.domain.form.ListElementDefinition>'
    )
    columns = catalog.grid_columns(form_xml_text)
    # elementDefinition 內也有 <id>，挖乾淨才不會取到錯的那一個
    assert columns == {'G_s_Sys_Name': '系統名稱'}


# ---------------------------------------------------------------- 查詢條件

def test_date_range_end_is_inclusive():
    start, end = query.date_range('2024-01-01', '2024-01-31')
    assert start == datetime.datetime(2024, 1, 1)
    assert end == datetime.datetime(2024, 2, 1)


def test_date_range_rejects_reversed_and_bad_format():
    with pytest.raises(query.QueryError):
        query.date_range('2024-02-01', '2024-01-01')
    with pytest.raises(query.QueryError):
        query.date_range('2024/01/01', '2024-01-31')
    with pytest.raises(query.QueryError):
        query.date_range('', '2024-01-31')


def test_state_name_marks_unknown_codes():
    assert query.state_name(query.PROCESS_STATE, 4) == '已作廢'
    # 沒實證的代碼不編名字，原樣標出來
    assert query.state_name(query.PROCESS_STATE, 99) == '狀態代碼 99'


def test_match_custom_is_and_and_case_insensitive():
    parsed = {'fields': {'a': 'EasyFlow GP', 'b': '資訊服務部'}}
    assert query.match_custom(parsed, [{'fieldId': 'a', 'value': 'easyflow'}])
    assert query.match_custom(parsed, [{'fieldId': 'a', 'value': 'GP'},
                                       {'fieldId': 'b', 'value': '資訊'}])
    assert not query.match_custom(parsed, [{'fieldId': 'a', 'value': 'GP'},
                                           {'fieldId': 'b', 'value': '人資'}])


def test_progress_gives_no_closed_time_while_pending():
    pending = [
        {'workItemName': '填表人', 'performerId': 'A', 'performerName': '甲',
         'createdTime': datetime.datetime(2024, 1, 1),
         'completedTime': datetime.datetime(2024, 1, 2)},
        {'workItemName': '主管', 'performerId': 'B', 'performerName': '乙',
         'createdTime': datetime.datetime(2024, 1, 2), 'completedTime': None},
    ]
    result = query.progress(pending)
    assert result['closedTime'] is None
    assert result['currentStep'] == '主管'
    assert result['completedSteps'] == 1


# ---------------------------------------------------------------- Excel

def _records():
    parsed = form_xml.parse(SAMPLE)
    return [{
        'summary': {
            'processSerialNumber': 'SIC003_00000001',
            'formSerialNumber': 'ITAccount_Application00000001',
            'processId': 'SIC003_', 'processName': '資訊帳號申請單',
            'subject': '測試', 'stateName': '已結案',
            'createdTime': datetime.datetime(2024, 1, 1, 9, 0, 0),
            'requesterId': 'S014060', 'requesterName': '王金平',
            'orgUnitId': 'S0610', 'orgUnitName': '資訊部',
            'abortComment': '', 'abortedBy': '', 'parseError': '',
        },
        'parsed': parsed,
        'items': [{'workItemName': '填表人', 'stateName': '已完成',
                   'performerId': 'S014060', 'performerName': '王金平',
                   'createdTime': datetime.datetime(2024, 1, 1, 9, 0, 0),
                   'completedTime': datetime.datetime(2024, 1, 1, 10, 0, 0),
                   'comment': '同意'}],
        'progress': {'closedTime': datetime.datetime(2024, 1, 1, 10, 0, 0),
                     'completedSteps': 1, 'totalSteps': 1,
                     'currentStep': '', 'currentPerformer': '',
                     'lastSigner': 'S014060 王金平',
                     'lastSignedTime': datetime.datetime(2024, 1, 1, 10, 0, 0)},
    }]


def _build(**overrides):
    options = {
        'list': True, 'content': True, 'signatures': True,
        'selected_fields': [('s_Req', '申請人代號'), ('Dlg', 'Dlg')],
        'grids': [{'id': 'Grid11', 'name': '帳號明細',
                   'columns': [{'id': 'G_a', 'name': '系統名稱', 'named': True},
                               {'id': 'G_b', 'name': '註銷/申請', 'named': True}]}],
        'catalog_fields': [{'id': 's_Req', 'name': '申請人代號', 'type': 'TEXTBOX',
                            'named': True, 'source': 'definition'},
                           {'id': 'Dlg', 'name': 'Dlg', 'type': 'DIALOGINPUT',
                            'named': False, 'source': 'definition'}],
        'meta': [('資料來源', '測試')],
    }
    options.update(overrides)
    return load_workbook(io.BytesIO(excel.build(_records(), options)))


def test_build_creates_one_sheet_per_section():
    workbook = _build()
    assert workbook.sheetnames == ['匯出條件', '清單', '內容', '明細-帳號明細',
                                   '欄位對照', '簽核名單']


def test_content_headers_use_field_names_and_detail_is_separate():
    workbook = _build()
    content = workbook['內容']
    assert [c.value for c in content[1]][-2:] == ['申請人代號', 'Dlg']
    detail = workbook['明細-帳號明細']
    assert [c.value for c in detail[1]] == ['流程單號', '表單單號', '明細列序',
                                            '系統名稱', '註銷/申請']
    # 兩筆明細各佔一列，不會被壓進單頭的儲存格
    assert detail.max_row == 3


def test_duplicate_field_names_get_id_suffix():
    workbook = _build(selected_fields=[('s_Req', '部門'), ('Dlg', '部門')])
    headers = [c.value for c in workbook['內容'][1]][-2:]
    assert headers == ['部門 (s_Req)', '部門 (Dlg)']


def test_unselected_sections_are_absent():
    workbook = _build(list=False, content=False)
    assert workbook.sheetnames == ['匯出條件', '簽核名單']


def test_illegal_control_characters_do_not_break_the_file():
    records = _records()
    records[0]['summary']['subject'] = 'a\x07b'
    content = excel.build(records, {'list': True, 'meta': []})
    assert load_workbook(io.BytesIO(content))['清單'].cell(row=2, column=5).value == 'ab'


# ---------------------------------------------------------------- SQL 結構

def test_form_join_uses_apply_so_rows_do_not_duplicate():
    """一個流程的 LocalRelevantData 有多列變數（實測 GeneralAffairs 每張單 3 列）。

    直接 LEFT JOIN 會讓同一張單重複出現，並多出表單單號空白的假列 ——
    這個 bug 真的發生過，靠 OUTER APPLY + TOP 1 修掉，這裡守住不要退回去。
    """
    sql = query.BASE_SQL
    assert 'OUTER APPLY' in sql
    assert 'SELECT TOP 1' in sql
    assert 'LEFT JOIN LocalRelevantData' not in sql
    # 必須真的對到 FormInstance，否則沒掛表單的變數列會被當成表單
    assert 'JOIN FormInstance f ON f.OID = l.valueOID' in sql


def test_form_columns_are_optional_to_avoid_reading_ntext():
    """沒要表單內容時不該把 fieldValues（ntext）撈出來。"""
    without = query.BASE_SQL.format(form_columns='', apply_columns='')
    assert 'fieldValues' not in without
    with_form = query.BASE_SQL.format(form_columns=query.FORM_COLUMN,
                                      apply_columns=query.APPLY_FORM_COLUMN)
    assert 'fieldValues' in with_form
